import os
import shutil
import tempfile
import unittest
from ngelaprak.core import (
    LaprakDocxBuilder,
    classify_file,
    infer_workspace_context,
    organize_workspace
)

TEST_OUT_DIR = os.path.join(os.path.dirname(__file__), "test_outputs")
os.makedirs(TEST_OUT_DIR, exist_ok=True)

class TestCore(unittest.TestCase):
    def test_docx_builder(self):
        builder = LaprakDocxBuilder()
        builder.add_section_heading("I. TUJUAN")
        builder.add_paragraph_text("1. Mahasiswa memahami struktur data.")
        
        out_docx = os.path.join(TEST_OUT_DIR, "sample_test.docx")
        res = builder.save(out_docx)
        self.assertTrue(os.path.exists(res))
        self.assertGreater(os.path.getsize(res), 500)

    def test_classify_file(self):
        self.assertEqual(classify_file("Modul 3.pdf"), "modul")
        self.assertEqual(classify_file("TP MODUL 3 STD.pdf"), "tp")
        self.assertEqual(classify_file("STRUKTUR-DATA_MOD-2_109092500017_SatriaBaktiWijaya.pdf"), "laprak")
        self.assertEqual(classify_file("main.cpp"), "code")
        self.assertEqual(classify_file("schema.sql"), "code")
        self.assertEqual(classify_file("screenshot.png"), "screenshot")
        self.assertEqual(classify_file(".gitignore"), "ignore")

    def test_organize_workspace_and_inference(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            # Create loose files
            f_mod = os.path.join(tmp_dir, "Modul 4.pdf")
            f_tp = os.path.join(tmp_dir, "TP MODUL 4.pdf")
            f_lap = os.path.join(tmp_dir, "STRUKTUR-DATA_MOD-3_109092500017_SatriaBaktiWijaya.docx")
            f_code = os.path.join(tmp_dir, "graph.cpp")
            
            with open(f_mod, "w") as f:
                f.write("dummy modul")
            with open(f_tp, "w") as f:
                f.write("dummy tp")
            with open(f_code, "w") as f:
                f.write("// graph implementation in C++\n#include <iostream>\n")

            builder = LaprakDocxBuilder()
            builder.add_section_heading("LAPORAN PRAKTIKUM")
            builder.save(f_lap)

            # Organize
            res = organize_workspace(tmp_dir)
            self.assertTrue(res["scaffolded"])
            self.assertEqual(len(res["moved_files"]), 4)

            # Check directory structure
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "Modul", "Modul 4.pdf")))
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "Modul", "TP MODUL 4.pdf")))
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "Laprak", "STRUKTUR-DATA_MOD-3_109092500017_SatriaBaktiWijaya.docx")))
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "Code", "graph.cpp")))
            self.assertTrue(os.path.exists(os.path.join(tmp_dir, "Laprak", "screenshots")))

            # Check context
            ctx = res["context"]
            self.assertEqual(ctx["target_module_number"], 4)
            self.assertEqual(ctx["language"], "C++")
            self.assertEqual(ctx["recommended_ide"], "codeblocks")
            self.assertTrue(ctx["has_code"])
            self.assertEqual(ctx["student_nim"], "109092500017")

    def test_docx_template_section_setters_and_typography(self):
        import docx
        tpl_path = os.path.join(os.path.dirname(__file__), "..", "ngelaprak", "templates", "template_kampus.docx")
        self.assertTrue(os.path.exists(tpl_path))

        builder = LaprakDocxBuilder(template_path=tpl_path)
        builder.set_cover_title("PERTEMUAN 5", "MODUL 5: TESTING JDBC")
        builder.set_student_info("Test Mahasiswa", "109092599999")
        builder.set_objectives([
            "1. Memahami konsep konektivitas database.",
            "2. Mengimplementasikan operasi query CRUD.",
        ])
        builder.set_tools("NetBeans 21, JDK 21, MySQL 8.0")
        builder.set_dasar_teori([
            "JDBC adalah standar antarmuka pemrograman aplikasi pada Java.",
            "Arsitektur MVC memisahkan logika akses data dari tampilan visual."
        ])
        builder.clear_guided_section(keep_until_p_index=35)
        builder.add_code_item("1.) Koneksi.java", [], "Penjelasan kelas koneksi database.")
        builder.add_output_section([])

        out_docx = os.path.join(TEST_OUT_DIR, "test_typography_result.docx")
        builder.save(out_docx)

        # Inspect resulting document
        doc = docx.Document(out_docx)
        secs = builder.find_section_indices()

        # 1. Verify Objectives are strictly non-bold
        start_obj = secs["TUJUAN"] + 1
        end_obj = secs["TOOL"]
        for p_idx in range(start_obj, end_obj):
            p = doc.paragraphs[p_idx]
            for r in p.runs:
                if r.text.strip():
                    self.assertFalse(r.bold, f"Objective run should not be bold: {r.text!r}")

        # 2. Verify Tool is strictly non-bold
        p_tool = doc.paragraphs[secs["TOOL"] + 1]
        for r in p_tool.runs:
            if r.text.strip():
                self.assertFalse(r.bold, f"Tool run should not be bold: {r.text!r}")

        # 3. Verify Dasar Teori paragraphs are strictly non-bold
        start_dt = secs["DASAR TEORI"] + 1
        end_dt = secs["GUIDED"]
        for p_idx in range(start_dt, end_dt):
            p = doc.paragraphs[p_idx]
            for r in p.runs:
                if r.text.strip():
                    self.assertFalse(r.bold, f"Dasar Teori run should not be bold: {r.text!r}")

        # 4. Verify main headings ARE bold
        for h_key in ["TUJUAN", "TOOL", "DASAR TEORI", "GUIDED"]:
            p_h = doc.paragraphs[secs[h_key]]
            for r in p_h.runs:
                if r.text.strip():
                    self.assertTrue(r.bold, f"Main heading {h_key} must be bold: {r.text!r}")

    def test_non_telkom_template_detection_and_builder(self):
        import docx
        from ngelaprak.core import analyze_template_docx
        
        # 1. Create a non-Telkom template docx (e.g. Institut Teknologi Sepuluh Nopember with table cover and NRP)
        tpl_path = os.path.join(TEST_OUT_DIR, "template_its_test.docx")
        doc = docx.Document()
        doc.add_paragraph("LAPORAN PRAKTIKUM PEMROGRAMAN WEB")
        doc.add_paragraph("MODUL 3: BACKEND REST API")
        
        # Add cover table (common in ITS / UI / UGM)
        table = doc.add_table(rows=3, cols=2)
        table.rows[0].cells[0].paragraphs[0].text = "Nama"
        table.rows[0].cells[1].paragraphs[0].text = ": Budi Prakoso"
        table.rows[1].cells[0].paragraphs[0].text = "NRP"
        table.rows[1].cells[1].paragraphs[0].text = ": 5025211001"
        table.rows[2].cells[0].paragraphs[0].text = "Kelas"
        table.rows[2].cells[1].paragraphs[0].text = ": Teknik Informatika A"
        
        doc.add_paragraph("DEPARTEMEN TEKNIK INFORMATIKA")
        doc.add_paragraph("FAKULTAS TEKNOLOGI ELEKTRO DAN INFORMATIKA CERDAS")
        doc.add_paragraph("INSTITUT TEKNOLOGI SEPULUH NOPEMBER")
        doc.add_paragraph("2026")
        doc.save(tpl_path)

        # 2. Test analyze_template_docx
        meta = analyze_template_docx(tpl_path)
        self.assertEqual(meta["student_name"], "Budi Prakoso")
        self.assertEqual(meta["student_nim"], "5025211001")
        self.assertEqual(meta["id_type"], "NRP")
        self.assertIn("INSTITUT TEKNOLOGI SEPULUH NOPEMBER", meta["kampus"].upper())
        self.assertIn("DEPARTEMEN TEKNIK INFORMATIKA", meta["prodi"].upper())
        self.assertIn("FAKULTAS TEKNOLOGI ELEKTRO", meta["fakultas"].upper())

        # 3. Test LaprakDocxBuilder with non-Telkom template
        builder = LaprakDocxBuilder(template_path=tpl_path)
        builder.set_student_info(
            nama="Siti Aminah",
            nim="5025211099",
            universitas="INSTITUT TEKNOLOGI SEPULUH NOPEMBER SURABAYA"
        )
        out_its = os.path.join(TEST_OUT_DIR, "its_result.docx")
        builder.save(out_its)

        # Verify updated values in table and paragraphs
        doc_res = docx.Document(out_its)
        res_text = "\n".join(p.text for p in doc_res.paragraphs)
        tbl_text = "\n".join(" ".join(c.text for c in r.cells) for t in doc_res.tables for r in t.rows)

        self.assertIn("Siti Aminah", tbl_text)
        self.assertIn("5025211099", tbl_text)
        self.assertIn("INSTITUT TEKNOLOGI SEPULUH NOPEMBER SURABAYA", res_text)

    def test_create_default_cover_and_logo(self):
        builder = LaprakDocxBuilder()
        builder.create_default_cover(
            title="LAPORAN PRAKTIKUM BASIS DATA",
            pertemuan="PERTEMUAN 2",
            modul_title="MODUL 2: DDL & DML",
            nama="Ahmad Dahlan",
            nim="21081010055",
            prodi="PROGRAM STUDI SISTEM INFORMASI",
            fakultas="FAKULTAS ILMU KOMPUTER",
            universitas="UNIVERSITAS INDONESIA",
            tahun="2026"
        )
        builder.add_main_heading("I. TUJUAN")
        builder.add_paragraph_text("1. Mahasiswa memahami query DDL.")
        
        out_def = os.path.join(TEST_OUT_DIR, "default_cover_test.docx")
        builder.save(out_def)
        self.assertTrue(os.path.exists(out_def))
        self.assertGreater(os.path.getsize(out_def), 1000)

if __name__ == "__main__":
    unittest.main()


