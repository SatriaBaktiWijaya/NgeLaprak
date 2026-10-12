import os
import re
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

class LaprakDocxBuilder:
    def __init__(self, template_path=None):
        self.template_path = template_path
        if template_path and os.path.exists(template_path):
            self.doc = docx.Document(template_path)
        else:
            self.doc = docx.Document()
            self._setup_default_styles()

    def _setup_default_styles(self):
        # Configure standard 1-inch margins
        for section in self.doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

    def find_section_indices(self):
        """
        Dynamically finds paragraph indices for major sections in the document:
        TUJUAN, TOOL, DASAR TEORI, GUIDED, KESIMPULAN.
        """
        sections = {}
        for idx, p in enumerate(self.doc.paragraphs):
            t = p.text.strip().upper()
            if "TUJUAN" in t and ("I." in t or t == "TUJUAN"):
                sections["TUJUAN"] = idx
            elif "TOOL" in t and ("II." in t or t == "TOOL"):
                sections["TOOL"] = idx
            elif "DASAR TEORI" in t and ("III." in t or t == "DASAR TEORI"):
                sections["DASAR TEORI"] = idx
            elif "GUIDED" in t and ("IV." in t or t == "GUIDED"):
                sections["GUIDED"] = idx
            elif "KESIMPULAN" in t and ("V." in t or t == "KESIMPULAN"):
                sections["KESIMPULAN"] = idx
        return sections

    def safe_set_paragraph_text(
        self,
        p,
        text,
        bold=False,
        font_name="Times New Roman",
        font_size=Pt(11),
        align=None,
        first_line_indent=None,
        left_indent=None,
        line_spacing=None,
        space_before=None,
        space_after=None,
        style=None
    ):
        """
        Safely sets paragraph text while strictly controlling typography and
        preventing unwanted bold inheritance from styles like Heading 1.
        """
        if style:
            p.style = style
        p.text = ""  # Clears all existing runs
        if align is not None:
            p.alignment = align
        if first_line_indent is not None:
            p.paragraph_format.first_line_indent = first_line_indent
        if left_indent is not None:
            p.paragraph_format.left_indent = left_indent
        if line_spacing is not None:
            p.paragraph_format.line_spacing = line_spacing
        if space_before is not None:
            p.paragraph_format.space_before = space_before
        if space_after is not None:
            p.paragraph_format.space_after = space_after

        r = p.add_run(text)
        r.font.name = font_name
        r.font.size = font_size
        r.font.bold = bold  # Explicitly False writes <w:b w:val="0"/> into XML
        return p

    def set_cover_title(self, pertemuan_text, modul_title_text):
        """Updates cover meeting header and module title."""
        if len(self.doc.paragraphs) > 2:
            self.doc.paragraphs[1].text = pertemuan_text
            self.doc.paragraphs[2].text = f"“{modul_title_text}”"

    def set_student_info(
        self,
        nama,
        nim,
        asisten_list=None,
        prodi="PROGRAM STUDI S1 REKAYASA PERANGKAT LUNAK",
        fakultas="FAKULTAS INFORMATIKA",
        universitas="TELKOM UNIVERSITY PURWOKERTO",
        tahun="2026"
    ):
        """
        Updates student identity on the cover page following exact template weights:
        - Nama : <Nama> (Bold 12pt)
        - Nama (NIM): <NIM> (Non-bold 12pt)
        - Asisten Praktikum: (Bold label, Non-bold assistant names)
        """
        for p in self.doc.paragraphs[:25]:
            t = p.text.strip()
            if t.startswith("Nama :") or t == "Nama":
                self.safe_set_paragraph_text(p, f"Nama : {nama}", bold=True, font_size=Pt(12))
            elif "Nama (NIM):" in t or "(NIM):" in t or "NIM :" in t:
                self.safe_set_paragraph_text(p, f"Nama (NIM): {nim}", bold=False, font_size=Pt(12))
            elif t.startswith("PROGRAM STUDI"):
                self.safe_set_paragraph_text(p, prodi, bold=True, font_size=Pt(14), align=WD_ALIGN_PARAGRAPH.CENTER)
            elif t.startswith("FAKULTAS") or "FAKULTAS" in t:
                self.safe_set_paragraph_text(p, fakultas, bold=True, font_size=Pt(14), align=WD_ALIGN_PARAGRAPH.CENTER)
            elif "TELKOM UNIVERSITY" in t:
                self.safe_set_paragraph_text(p, universitas, bold=True, font_size=Pt(14), align=WD_ALIGN_PARAGRAPH.CENTER)
            elif t.isdigit() and len(t) == 4:
                self.safe_set_paragraph_text(p, tahun, bold=True, font_size=Pt(14), align=WD_ALIGN_PARAGRAPH.CENTER)

        if asisten_list:
            asist_found = False
            asist_idx = 0
            for p in self.doc.paragraphs[:25]:
                t = p.text.strip()
                if "Asisten Praktikum" in t:
                    asist_found = True
                    continue
                if asist_found and asist_idx < len(asisten_list) and "(" in t and ")" in t:
                    self.safe_set_paragraph_text(p, asisten_list[asist_idx], bold=False, font_size=Pt(12))
                    asist_idx += 1

    def set_objectives(self, objectives, font_size=Pt(12), line_spacing=1.15):
        """
        Sets the TUJUAN (Objectives) items.
        Strictly NON-BOLD, Times New Roman, 12pt.
        """
        if isinstance(objectives, str):
            objectives = [line.strip() for line in objectives.splitlines() if line.strip()]

        secs = self.find_section_indices()
        if "TUJUAN" in secs and "TOOL" in secs:
            start = secs["TUJUAN"] + 1
            end = secs["TOOL"]
            existing_count = end - start

            for i, obj in enumerate(objectives):
                text = obj if re.match(r"^\d+\.\s*", obj) else f"{i + 1}. {obj}"
                if i < existing_count:
                    p = self.doc.paragraphs[start + i]
                    self.safe_set_paragraph_text(
                        p, text, bold=False, font_size=font_size, line_spacing=line_spacing, space_after=Pt(2), style='Normal'
                    )
                else:
                    p_tool = self.doc.paragraphs[secs["TOOL"]]
                    p_new = p_tool.insert_paragraph_before(style='Normal')
                    self.safe_set_paragraph_text(
                        p_new, text, bold=False, font_size=font_size, line_spacing=line_spacing, space_after=Pt(2)
                    )
            if len(objectives) < existing_count:
                for _ in range(existing_count - len(objectives)):
                    p_to_del = self.doc.paragraphs[start + len(objectives)]
                    p_to_del._element.getparent().remove(p_to_del._element)
        else:
            self.add_main_heading("I. TUJUAN")
            for i, obj in enumerate(objectives):
                text = obj if re.match(r"^\d+\.\s*", obj) else f"{i + 1}. {obj}"
                self.add_paragraph_text(text, bold=False, font_size=font_size)

    def set_tools(self, tools_text, font_size=Pt(11), left_indent=Inches(0.4)):
        """
        Sets the TOOL description text.
        Strictly NON-BOLD, Times New Roman, 11pt, indented left.
        """
        secs = self.find_section_indices()
        if "TOOL" in secs:
            target_idx = secs["TOOL"] + 1
            if target_idx < len(self.doc.paragraphs):
                p = self.doc.paragraphs[target_idx]
                self.safe_set_paragraph_text(
                    p, tools_text, bold=False, font_size=font_size, left_indent=left_indent, line_spacing=1.15, style='Normal'
                )
                return p
        self.add_main_heading("II. TOOL")
        p = self.add_paragraph_text(tools_text, bold=False, font_size=font_size)
        p.paragraph_format.left_indent = left_indent
        return p

    def set_dasar_teori(self, text_or_paragraphs, font_size=Pt(12), first_line_indent=Inches(0.35), line_spacing=1.15):
        """
        Sets the DASAR TEORI explanation paragraphs.
        Strictly NON-BOLD, Times New Roman, 12pt, Justified, First-Line Indented.
        """
        if isinstance(text_or_paragraphs, str):
            paras = [p.strip() for p in text_or_paragraphs.split("\n\n") if p.strip()]
            if not paras:
                paras = [text_or_paragraphs.strip()]
        else:
            paras = list(text_or_paragraphs)

        secs = self.find_section_indices()
        if "DASAR TEORI" in secs and "GUIDED" in secs:
            start = secs["DASAR TEORI"] + 1
            end = secs["GUIDED"]
            existing_count = end - start
            for i, para_text in enumerate(paras):
                if i < existing_count:
                    p = self.doc.paragraphs[start + i]
                    self.safe_set_paragraph_text(
                        p,
                        para_text,
                        bold=False,
                        font_size=font_size,
                        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                        first_line_indent=first_line_indent,
                        line_spacing=line_spacing,
                        space_after=Pt(6),
                        style='Normal'
                    )
                else:
                    p_guided = self.doc.paragraphs[secs["GUIDED"]]
                    p_new = p_guided.insert_paragraph_before(style='Normal')
                    self.safe_set_paragraph_text(
                        p_new,
                        para_text,
                        bold=False,
                        font_size=font_size,
                        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                        first_line_indent=first_line_indent,
                        line_spacing=line_spacing,
                        space_after=Pt(6)
                    )
            if len(paras) < existing_count:
                for _ in range(existing_count - len(paras)):
                    p_to_del = self.doc.paragraphs[start + len(paras)]
                    p_to_del._element.getparent().remove(p_to_del._element)
        else:
            self.add_main_heading("III. DASAR TEORI")
            for para_text in paras:
                self.add_explanation_paragraph(para_text, indent=first_line_indent, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    def clear_guided_section(self, keep_until_p_index=35):
        """Clears existing guided paragraphs in a template to rebuild fresh."""
        while len(self.doc.paragraphs) > keep_until_p_index:
            p = self.doc.paragraphs[-1]
            p._element.getparent().remove(p._element)

    def add_main_heading(self, text):
        """
        Adds a top-level Roman numeral section heading (e.g. 'I. TUJUAN', 'IV. GUIDED').
        Bold, 14pt, Times New Roman.
        """
        p = self.doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        return p

    def add_section_heading(self, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """
        Adds a subsection heading (e.g. '1.) TestJDBC.java').
        Default NON-BOLD matching reference template.
        """
        p = self.doc.add_paragraph(style='Normal')
        p.alignment = align
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = bold
        return p

    def add_label(self, text, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """
        Adds a section label (e.g. 'Kode', 'Penjelasan').
        Default NON-BOLD matching reference template.
        """
        p = self.doc.add_paragraph(style='Normal')
        p.alignment = align
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = bold
        return p

    def add_explanation_paragraph(self, text, indent=Inches(0.4), align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """
        Adds an indented explanation paragraph with normal non-bold Times New Roman font.
        """
        p = self.doc.add_paragraph(style='Normal')
        p.alignment = align
        p.paragraph_format.first_line_indent = indent
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = False
        return p

    def add_paragraph_text(self, text, bold=False, font_size=Pt(11), indent=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = self.doc.add_paragraph(style='Normal')
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = indent
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = font_size
        r.font.bold = bold
        return p

    def add_code_item(self, item_title, image_paths, explanation_text, bold_title=False, bold_labels=False, img_width=Inches(5.2)):
        """
        Appends a code task item:
        - Title (e.g. '1.) TestJDBC.java') -> non-bold
        - 'Kode' -> non-bold
        - Embedded screenshots
        - 'Penjelasan' -> non-bold
        - Indented non-bold explanation paragraph
        """
        self.add_section_heading(item_title, bold=bold_title)
        self.add_label("Kode", bold=bold_labels)

        w = img_width if img_width is not None else Inches(5.2)
        for img_path in image_paths:
            if os.path.exists(img_path):
                p_img = self.doc.add_paragraph(style='Normal')
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(4)
                p_img.paragraph_format.space_after = Pt(6)
                p_img.add_run().add_picture(img_path, width=w)

        self.add_label("Penjelasan", bold=bold_labels)
        self.add_explanation_paragraph(explanation_text)

        # Subtle blank space
        p_blank = self.doc.add_paragraph(style='Normal')
        p_blank.paragraph_format.space_before = Pt(2)
        p_blank.paragraph_format.space_after = Pt(2)

    def add_output_section(self, output_images, default_width=Inches(4.8), bold_heading=False):
        """
        Appends the Output section with screenshots.
        'Output' label is non-bold by default matching reference template.
        """
        self.add_section_heading("Output", bold=bold_heading)
        for item in output_images:
            if isinstance(item, tuple):
                img_path, w = item
            else:
                img_path, w = item, default_width

            if os.path.exists(img_path):
                p_img = self.doc.add_paragraph(style='Normal')
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(4)
                p_img.paragraph_format.space_after = Pt(6)
                p_img.add_run().add_picture(img_path, width=w)

    def sanitize_document_typography(self):
        """
        Scans all paragraphs in the document to enforce strict typography standards:
        1. Roman numeral main section headings ('I. TUJUAN', 'II. TOOL', 'III. DASAR TEORI',
           'IV. GUIDED', etc.) are BOLD (14pt).
        2. All body paragraphs (Objectives, Tools, Dasar Teori explanations, Guided code
           items, labels, and code explanations) are strictly NON-BOLD (bold = False).
        3. Fonts are standardized to Times New Roman.
        """
        main_keywords = {"TUJUAN", "TOOL", "DASAR TEORI", "GUIDED", "KESIMPULAN", "TUGAS", "PENDAHULUAN"}

        for idx, p in enumerate(self.doc.paragraphs):
            # Preserve cover page formatting (usually first 24 paragraphs)
            if idx < 24:
                continue

            t = p.text.strip()
            if not t:
                continue

            # Detect if paragraph is a main section heading
            has_num_pr = p._p.pPr is not None and p._p.pPr.find(docx.oxml.ns.qn('w:numPr')) is not None
            is_roman_header = bool(re.match(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X)\.\s+", t))
            is_keyword_header = has_num_pr and t.upper() in main_keywords

            is_main_heading = is_roman_header or is_keyword_header

            if is_main_heading:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.bold = True
                    r.font.size = Pt(14)
            else:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.bold = False

    def save(self, output_path, sanitize=True):
        """
        Saves document, automatically sanitizing typography if requested.
        """
        if sanitize:
            self.sanitize_document_typography()
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        self.doc.save(output_path)
        return output_path
